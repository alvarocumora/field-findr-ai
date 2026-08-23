"""
Láminas planimétricas periciales OT-PER-2026-0087-V5 (P-01 y P-02).

Entradas (todas preexistentes, sin geometría inventada):
  - /tmp/anexo4b/work/parcelas_v4b.json  (reconstrucción GIS V5: 754 lotes, faja, 59 torres)
  - /mnt/user-uploads/KMZ_Sol_de_Caone_Rev0-4.kmz  (eje LT, caminos de acceso)
  - /mnt/documents/Base_Empirica_V5_LT220_Caone_Mataquito.xlsx  (IDV y clase V0-V4 del V5)
  - Cuadro PES oficial notificado (14 lotes, 35.777 m²) — prevalece para efectos jurídicos.

Salidas: PDF A3 apaisado vectorial, PNG 300 dpi y tabla XLSX/CSV/GeoJSON.
"""
import csv
import json
import os
import re
import zipfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch, Polygon as MplPoly, Rectangle
from openpyxl import Workbook, load_workbook
from pyproj import Transformer
from shapely import wkt as shwkt
from shapely.geometry import LineString, MultiLineString, Polygon, mapping
from shapely.ops import unary_union

WORK = "/tmp/anexo4b/work"
OUT = "/mnt/documents"
UP = "/mnt/user-uploads"
LT_KMZ = f"{UP}/KMZ_Sol_de_Caone_Rev0-4.kmz"
XLSX = f"{OUT}/Base_Empirica_V5_LT220_Caone_Mataquito.xlsx"
FOLIO = "OT-PER-2026-0087-V5"
T = Transformer.from_crs("EPSG:4326", "EPSG:32719", always_xy=True)

# --- Cuadro PES oficial notificado (dato jurídico, no recalculado) ----------
PES = {  # lote: (superficie lote m2, superficie PES m2, % afectación)
    61: (6500, 5574, 85.8), 60: (5500, 3710, 67.5), 57: (6000, 3572, 59.5),
    56: (5500, 3130, 56.9), 55: (5500, 2817, 51.2), 54: (6300, 2965, 47.1),
    53: (5000, 2207, 44.1), 52: (5000, 2057, 41.1), 51: (5000, 1967, 39.3),
    50: (5100, 1904, 37.3), 49: (5000, 1951, 39.0), 48: (5000, 1857, 37.1),
    47: (6000, 1822, 30.4), 62: (5300, 244, 4.6),
}
PES_TOTAL, SUP_TOTAL, PCT_TOTAL = 35777, 76700, 46.6
LARGO_FAJA, ANCHO_FAJA = 1198.33, 29.86
CLASES = ["V0", "V1", "V2", "V3", "V4"]
CLASE_TOTALES = {"V0": 132, "V1": 194, "V2": 338, "V3": 60, "V4": 30}
FILL = {"V0": "#f2f2f2", "V1": "#f0e6cf", "V2": "#f7e08a", "V3": "#f2b06a", "V4": "#e2603c"}
CLASE_TXT = {
    "V0": "exposición nula / muy baja", "V1": "exposición baja",
    "V2": "exposición media", "V3": "exposición alta", "V4": "exposición muy alta",
}
RECL = "Rol 168-915"
FONT = "DejaVu Sans"
plt.rcParams.update({"font.family": FONT, "pdf.fonttype": 42, "ps.fonttype": 42})


# ------------------------------------------------------------------ datos --
def load_caminos():
    with zipfile.ZipFile(LT_KMZ) as z:
        name = [f for f in z.namelist() if f.endswith(".kml")][0]
        kml = z.read(name).decode("utf-8", "replace")
    lines = []
    for fname, body in re.findall(r"<Folder>\s*<name>(.*?)</name>(.*?)</Folder>", kml, re.S):
        if fname not in ("Acceso LTE", "Caminos Internos con Línea Eléctrica"):
            continue
        for c in re.finditer(r"<(LineString|Polygon)>(.*?)</\1>", body, re.S):
            cc = re.search(r"<coordinates>(.*?)</coordinates>", c.group(2), re.S)
            if not cc:
                continue
            pts = []
            for tok in cc.group(1).split():
                p = tok.split(",")
                if len(p) >= 2:
                    pts.append(T.transform(float(p[0]), float(p[1])))
            if len(pts) >= 2:
                lines.append(LineString(pts))
    return lines


def load_base():
    ws = load_workbook(XLSX)["BASE_EMPIRICA_754"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return {(r[0], int(r[1])): dict(
        sup=r[2], x=r[3], y=r[4], d_faja=r[5], d_torre=r[6], d_cond=r[7],
        idv=r[10], clase=r[11], cruzado=r[15], sup_faja=r[16]) for r in rows}


def load_geo():
    d = json.load(open(f"{WORK}/parcelas_v4b.json"))
    faja = shwkt.loads(d["faja_wkt"])
    torres = [tuple(t) for t in d["torres"]]
    lotes = []
    for rol, arr in d["predios"].items():
        for x in arr:
            lotes.append(dict(rol=rol, n=int(x["n"]), poly=shwkt.loads(x["wkt"]),
                              area=x["area_gis"], area_faja=x["area_faja"],
                              cruzada=x["cruzada"], d_faja=x["d_faja"], d_torre=x["d_torre"]))
    return lotes, faja, torres


# ------------------------------------------------------------ validaciones --
def validar(lotes, base, torres):
    errs = []
    if sum(v[1] for v in PES.values()) != PES_TOTAL:
        errs.append(f"Suma PES = {sum(v[1] for v in PES.values())} ≠ {PES_TOTAL}")
    if sum(v[0] for v in PES.values()) != SUP_TOTAL:
        errs.append(f"Suma superficies 14 lotes = {sum(v[0] for v in PES.values())} ≠ {SUP_TOTAL}")
    if round(PES_TOTAL / SUP_TOTAL * 100, 1) != PCT_TOTAL:
        errs.append("Porcentaje de afectación directa no reproduce 46,6 %")
    if len(PES) != 14:
        errs.append("El cuadro PES no contiene 14 lotes")
    if len(lotes) != 754:
        errs.append(f"Universo GIS = {len(lotes)} ≠ 754 lotes")
    ids = [(l["rol"], l["n"]) for l in lotes]
    if len(set(ids)) != len(ids):
        errs.append("Existen identificadores de lote duplicados")
    if any(l["poly"].is_empty or l["n"] is None for l in lotes):
        errs.append("Existen polígonos sin ID o vacíos")
    faltan = [i for i in ids if i not in base]
    if faltan:
        errs.append(f"{len(faltan)} lotes sin registro en la Base Empírica V5")
    cnt = {c: 0 for c in CLASES}
    for i in ids:
        cnt[base[i]["clase"]] += 1
    if cnt != CLASE_TOTALES:
        errs.append(f"Distribución de clases {cnt} ≠ {CLASE_TOTALES}")
    if sum(cnt.values()) != 754:
        errs.append("V0+V1+V2+V3+V4 ≠ 754")
    if len(torres) != 59 or any(len(t) != 2 for t in torres):
        errs.append(f"Estructuras georreferenciadas = {len(torres)} ≠ 59")
    for n in PES:
        if (RECL, n) not in base:
            errs.append(f"Lote notificado {n} sin geometría en el rol {RECL}")
    return errs, cnt


# --------------------------------------------------------------- dibujado --
def escala_real(ax_w_in, ancho_m):
    return ancho_m / (ax_w_in * 0.0254)


def barra_escala(ax, x0, y0, ancho_m):
    """Barra de escala gráfica en unidades de terreno."""
    for cand in (100, 200, 250, 500, 1000, 2000):
        if ancho_m / cand <= 6:
            paso = cand
            break
    else:
        paso = 2000
    for i in range(4):
        ax.add_patch(Rectangle((x0 + i * paso, y0), paso, ancho_m * 0.006,
                               facecolor="black" if i % 2 == 0 else "white",
                               edgecolor="black", lw=0.5, zorder=8))
    for i in range(5):
        ax.text(x0 + i * paso, y0 - ancho_m * 0.012, f"{i * paso:,}".replace(",", "."),
                ha="center", va="top", fontsize=5.5, zorder=8)
    ax.text(x0 + 2 * paso, y0 + ancho_m * 0.016, "metros", ha="center", fontsize=5.5, zorder=8)


def norte(ax, x, y, h):
    ax.annotate("", xy=(x, y + h), xytext=(x, y),
                arrowprops=dict(arrowstyle="-|>", color="black", lw=0.9), zorder=8)
    ax.text(x, y + h * 1.05, "N", ha="center", va="bottom", fontsize=8, fontweight="bold", zorder=8)


def dibujar_mapa(ax, lotes, base, faja, torres, caminos, bbox, etiquetas=None,
                 label_pct=False, grid=True):
    x0, y0, x1, y1 = bbox
    for l in lotes:
        p = l["poly"]
        if p.bounds[2] < x0 or p.bounds[0] > x1 or p.bounds[3] < y0 or p.bounds[1] > y1:
            continue
        clase = base[(l["rol"], l["n"])]["clase"]
        ax.add_patch(MplPoly(list(p.exterior.coords), closed=True,
                             facecolor=FILL[clase], alpha=0.45,
                             edgecolor="#555555", lw=0.25, zorder=2))
    # límite predial de cada rol
    for rol in sorted({l["rol"] for l in lotes}):
        u = unary_union([l["poly"] for l in lotes if l["rol"] == rol])
        geoms = u.geoms if u.geom_type == "MultiPolygon" else [u]
        for g in geoms:
            ax.plot(*g.exterior.xy, color="black", lw=0.9, zorder=4)
    # faja PES (hachurado)
    fg = faja.geoms if faja.geom_type == "MultiPolygon" else [faja]
    for g in fg:
        ax.add_patch(MplPoly(list(g.exterior.coords), closed=True, facecolor="none",
                             edgecolor="#b00020", lw=0.7, hatch="////", zorder=5))
        ax.add_patch(MplPoly(list(g.exterior.coords), closed=True, facecolor="#b00020",
                             alpha=0.10, edgecolor="none", zorder=3))
    # caminos
    for c in caminos:
        ax.plot(*c.xy, color="#7a6a55", lw=0.6, zorder=4)
    # eje LT + torres
    eje = LineString(sorted(torres, key=lambda t: (t[0], t[1])))
    ax.plot(*eje.xy, color="#8b005d", lw=0.8, ls=(0, (6, 3)), zorder=6)
    tx = [t[0] for t in torres]
    ty = [t[1] for t in torres]
    ax.scatter(tx, ty, s=11, facecolor="white", edgecolor="#8b005d", lw=0.8, zorder=7)
    if etiquetas:
        for l in lotes:
            if (l["rol"], l["n"]) not in etiquetas:
                continue
            c = l["poly"].representative_point()
            txt = str(l["n"])
            if label_pct:
                txt = f"{l['n']}\n{PES[l['n']][2]:.1f}%".replace(".", ",")
            ax.text(c.x, c.y, txt, ha="center", va="center", fontsize=5.2,
                    color="black", zorder=9, linespacing=1.1)
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect("equal")
    if grid:
        ax.tick_params(labelsize=5.5, direction="in", length=2.5)
        ax.ticklabel_format(style="plain", useOffset=False)
        ax.grid(True, color="#c8c8c8", lw=0.3, ls=(0, (2, 3)))
        for lab in ax.get_yticklabels():
            lab.set_rotation(90)
            lab.set_va("center")
    for s in ax.spines.values():
        s.set_linewidth(0.8)
        s.set_color("black")


# ------------------------------------------------------------ panel lateral --
def caja(ax, y, h, titulo):
    ax.add_patch(Rectangle((0, y - h), 1, h, transform=ax.transAxes, facecolor="none",
                           edgecolor="black", lw=0.7, clip_on=False))
    ax.add_patch(Rectangle((0, y - 0.021), 1, 0.021, transform=ax.transAxes,
                           facecolor="#e8e8e8", edgecolor="black", lw=0.7, clip_on=False))
    ax.text(0.012, y - 0.0105, titulo, transform=ax.transAxes, fontsize=6.6,
            fontweight="bold", va="center")
    return y - 0.028


def kv(ax, y, k, v, dy=0.0158, bold=False):
    ax.text(0.02, y, k, transform=ax.transAxes, fontsize=5.9, va="top",
            fontweight="bold" if bold else "normal")
    ax.text(0.98, y, v, transform=ax.transAxes, fontsize=5.9, va="top", ha="right",
            fontweight="bold" if bold else "normal")
    return y - dy


def parrafo(ax, y, txt, dy=0.0125, size=5.2):
    ax.text(0.02, y, txt, transform=ax.transAxes, fontsize=size, va="top", ha="left",
            wrap=False, linespacing=1.35)
    return y - dy * (txt.count("\n") + 1)


def cartucho(ax, y, lamina, escala_txt):
    y = caja(ax, y, 0.185, "CARTUCHO TÉCNICO")
    campos = [
        ("PROYECTO", "LT 1x220 kV Caone–Mataquito"),
        ("MANDANTE", "Forestal Valle Escondido SpA"),
        ("DOCUMENTO", "Plano Pericial de Afectación"),
        ("", "Directa y Visual"),
        ("LÁMINA", lamina),
        ("FOLIO", f"{FOLIO}-{lamina}"),
        ("SISTEMA", "UTM 19S – EPSG:32719"),
        ("ESCALA", escala_txt),
        ("FECHA", "agosto 2026"),
        ("FUENTE", "PES notificados / KMZ concesionaria /"),
        ("", "cartografía parcelaria / Base Empírica V5"),
        ("ELABORADO POR", "OfertaTerrenos · Valuation"),
        ("ESTADO", "Documento técnico para análisis pericial"),
    ]
    for k, v in campos:
        ax.text(0.02, y, k, transform=ax.transAxes, fontsize=5.2, va="top", color="#333333")
        ax.text(0.36, y, v, transform=ax.transAxes, fontsize=5.4, va="top")
        y -= 0.0125
    return y


def simbologia(ax, y, con_clases=True):
    h = 0.205 if con_clases else 0.115
    y = caja(ax, y, h, "SIMBOLOGÍA")
    items = [
        ("predial", "Límite predial (rol)"),
        ("lote", "Límite de lote"),
        ("faja", "Faja de servidumbre PES (afectación física)"),
        ("eje", "Eje LT 1x220 kV Caone–Mataquito"),
        ("torre", "Torre / estructura georreferenciada"),
        ("camino", "Camino de acceso"),
    ]
    for kind, label in items:
        yy = y - 0.004
        if kind == "predial":
            ax.plot([0.03, 0.10], [yy, yy], transform=ax.transAxes, color="black", lw=1.1)
        elif kind == "lote":
            ax.plot([0.03, 0.10], [yy, yy], transform=ax.transAxes, color="#555555", lw=0.5)
        elif kind == "faja":
            ax.add_patch(Rectangle((0.03, yy - 0.0045), 0.07, 0.009, transform=ax.transAxes,
                                   facecolor="none", edgecolor="#b00020", lw=0.7, hatch="////"))
        elif kind == "eje":
            ax.plot([0.03, 0.10], [yy, yy], transform=ax.transAxes, color="#8b005d",
                    lw=0.9, ls=(0, (5, 3)))
        elif kind == "torre":
            ax.plot([0.065], [yy], transform=ax.transAxes, marker="o", ms=3.2,
                    mfc="white", mec="#8b005d", mew=0.9, ls="none")
        else:
            ax.plot([0.03, 0.10], [yy, yy], transform=ax.transAxes, color="#7a6a55", lw=0.7)
        ax.text(0.13, y, label, transform=ax.transAxes, fontsize=5.4, va="top")
        y -= 0.0135
    if con_clases:
        y -= 0.004
        ax.text(0.02, y, "CLASE DE EXPOSICIÓN / DOMINANCIA VISUAL (modelo V5)",
                transform=ax.transAxes, fontsize=5.2, fontweight="bold", va="top")
        y -= 0.014
        for c in CLASES:
            ax.add_patch(Rectangle((0.03, y - 0.0085), 0.07, 0.009, transform=ax.transAxes,
                                   facecolor=FILL[c], alpha=0.45, edgecolor="#555555", lw=0.4))
            ax.text(0.13, y, f"{c} – {CLASE_TXT[c]}", transform=ax.transAxes,
                    fontsize=5.4, va="top")
            y -= 0.0135
        y -= 0.002
        ax.text(0.02, y, "V0–V4 expresa clase de exposición visual; no constituye\n"
                         "por sí sola un porcentaje de depreciación económica.",
                transform=ax.transAxes, fontsize=4.9, va="top", color="#333333", linespacing=1.3)
    return y - 0.03


def marco(fig, titulo, subt):
    fig.patches.append(Rectangle((0.014, 0.014), 0.972, 0.972, transform=fig.transFigure,
                                 facecolor="none", edgecolor="black", lw=1.1))
    fig.patches.append(Rectangle((0.021, 0.021), 0.958, 0.958, transform=fig.transFigure,
                                 facecolor="none", edgecolor="#777777", lw=0.5))
    fig.text(0.032, 0.958, titulo, fontsize=13, fontweight="bold", va="center")
    fig.text(0.032, 0.934, subt, fontsize=7.2, va="center", color="#222222", linespacing=1.4)


# ------------------------------------------------------------------ láminas --
def lamina_p01(lotes, base, faja, torres, caminos, cnt):
    fig = plt.figure(figsize=(16.54, 11.69))
    marco(fig, "PLANO PERICIAL DE AFECTACIÓN DIRECTA Y VISUAL",
          "Proyecto Valle Escondido / Fundo Hualañé  ·  Línea Eléctrica 1x220 kV Caone–Mataquito\n"
          "Hualañé – Región del Maule  ·  Folio " + FOLIO + "-P01")

    ax = fig.add_axes([0.032, 0.045, 0.615, 0.855])
    recl = unary_union([l["poly"] for l in lotes if l["rol"] == RECL])
    fx0, fy0, fx1, fy1 = unary_union([recl, faja]).bounds
    cx, cy = (fx0 + fx1) / 2, (fy0 + fy1) / 2
    w = max(fx1 - fx0, fy1 - fy0) * 1.45
    ratio = 0.615 * 16.54 / (0.855 * 11.69)
    half_x, half_y = w * ratio / 2, w / 2
    bbox = (cx - half_x, cy - half_y, cx + half_x, cy + half_y)
    dibujar_mapa(ax, lotes, base, faja, torres, caminos, bbox,
                 etiquetas={(RECL, n) for n in PES})
    ancho_m = bbox[2] - bbox[0]
    esc = escala_real(0.615 * 16.54, ancho_m)
    barra_escala(ax, bbox[0] + ancho_m * 0.04, bbox[1] + ancho_m * 0.045, ancho_m)
    norte(ax, bbox[2] - ancho_m * 0.05, bbox[1] + ancho_m * 0.04, ancho_m * 0.06)
    ax.set_xlabel("Este UTM 19S (m) – EPSG:32719", fontsize=5.8)
    ax.set_ylabel("Norte UTM 19S (m) – EPSG:32719", fontsize=5.8)
    esc_txt = f"1:{round(esc / 100) * 100:,}".replace(",", ".") + " (A3)"

    p = fig.add_axes([0.663, 0.045, 0.305, 0.855])
    p.axis("off")
    y = 1.0
    y = caja(p, y, 0.135, "CUADRO RESUMEN – AFECTACIÓN DIRECTA (PES OFICIAL)")
    y = kv(p, y, "Superficie 14 lotes notificados", "76.700 m²")
    y = kv(p, y, "Superficie PES notificada", "35.777 m²")
    y = kv(p, y, "Afectación directa", "46,6 %", bold=True)
    y = kv(p, y, "Longitud aproximada afectada", "1.198,33 m")
    y = kv(p, y, "Ancho medio faja PES", "aprox. 29,86 m")
    y = kv(p, y, "N° de lotes notificados", "14")
    y = kv(p, y, "Tensión", "220 kV")
    y -= 0.014

    y = caja(p, y, 0.152, "LOTES NOTIFICADOS – SUPERFICIE / PES / % AFECTACIÓN")
    p.text(0.02, y, "Lote", transform=p.transAxes, fontsize=5.2, fontweight="bold", va="top")
    p.text(0.30, y, "Sup. m²", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    p.text(0.55, y, "PES m²", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    p.text(0.78, y, "% afect.", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    p.text(0.98, y, "Clase V5", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    y -= 0.0115
    for n, (s, pes, pct) in sorted(PES.items(), key=lambda kv2: -kv2[1][2]):
        cl = base[(RECL, n)]["clase"]
        p.text(0.02, y, str(n), transform=p.transAxes, fontsize=5.2, va="top")
        p.text(0.30, y, f"{s:,}".replace(",", "."), transform=p.transAxes, fontsize=5.2,
               va="top", ha="right")
        p.text(0.55, y, f"{pes:,}".replace(",", "."), transform=p.transAxes, fontsize=5.2,
               va="top", ha="right")
        p.text(0.78, y, f"{pct:.1f} %".replace(".", ","), transform=p.transAxes,
               fontsize=5.2, va="top", ha="right")
        p.text(0.98, y, cl, transform=p.transAxes, fontsize=5.2, va="top", ha="right")
        y -= 0.0092
    p.plot([0.02, 0.98], [y + 0.0055, y + 0.0055], transform=p.transAxes, color="black", lw=0.5)
    y -= 0.001
    p.text(0.02, y, "TOTAL (14 lotes)", transform=p.transAxes, fontsize=5.2,
           fontweight="bold", va="top")
    p.text(0.30, y, "76.700", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    p.text(0.55, y, "35.777", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    p.text(0.78, y, "46,6 %", transform=p.transAxes, fontsize=5.2, fontweight="bold",
           va="top", ha="right")
    y -= 0.024

    y = caja(p, y, 0.086, "CONTROL DEL UNIVERSO VISUAL V5")
    y = kv(p, y, "Planos / roles analizados", "9", dy=0.0125)
    y = kv(p, y, "Lotes analizados", "754", dy=0.0125)
    y = kv(p, y, "Superficie total analizada", "433,7 ha", dy=0.0125)
    y = kv(p, y, "Estructuras georreferenciadas", "59", dy=0.0125)
    y = kv(p, y, "V0 / V1 / V2 / V3 / V4",
           f"{cnt['V0']} / {cnt['V1']} / {cnt['V2']} / {cnt['V3']} / {cnt['V4']}", dy=0.0125)
    y = kv(p, y, "Suma de clases (control)",
           f"{sum(cnt.values())} = 754 ✓", dy=0.0125, bold=True)
    y -= 0.012

    y = simbologia(p, y)

    y = caja(p, y, 0.118, "OBSERVACIONES TÉCNICAS")
    obs = [
        "1. La lámina integra la afectación física de los Planos Especiales de\n"
        "    Servidumbre con el análisis de exposición visual desarrollado en V5.",
        "2. Para efectos de superficie jurídicamente afectada prevalecen los PES\n"
        "    oficiales notificados (35.777 m² sobre 14 lotes).",
        "3. La clasificación V0–V4 corresponde al modelo de dominancia/exposición\n"
        "    visual V5 y no constituye por sí sola un porcentaje de depreciación.",
        "4. La infraestructura aérea genera una zona de influencia visual que excede\n"
        "    físicamente la superficie comprendida en la servidumbre.",
        "5. La eventual depreciación comercial debe determinarse mediante evidencia\n"
        "    de mercado y análisis pericial separado.",
    ]
    for t in obs:
        y = parrafo(p, y, t, dy=0.0105)
    y -= 0.012

    y = caja(p, y, 0.062, "NOTA DE PREVALENCIA Y TRAZABILIDAD")
    y = parrafo(p, y,
                "Para efectos indemnizatorios prevalecen los Planos Especiales de\n"
                "Servidumbre notificados. La reconstrucción GIS V5 se utiliza para\n"
                "análisis espacial y visual. La reconstrucción GIS identifica 13 cruces\n"
                "geométricos principales; la pretensión jurídica conserva los 14 lotes y\n"
                "35.777 m² de los PES (el lote 62 registra 244 m² de faja en el PES, bajo\n"
                "la tolerancia de la reconstrucción GIS). Diferencia informada expresamente.",
                dy=0.0105)
    y -= 0.010

    y = caja(p, y, 0.040, "CONTEXTO TERRITORIAL (NOTA SECUNDARIA)")
    y = parrafo(p, y,
                "Se ha identificado la existencia de oposición local organizada respecto\n"
                "de infraestructura de transmisión en el sector. Este antecedente se\n"
                "incorpora sólo como contexto territorial y no constituye por sí mismo\n"
                "prueba de depreciación económica. No se le asigna valor monetario.",
                dy=0.0105)
    y -= 0.012
    cartucho(p, y, "P01", esc_txt)

    fig.savefig(f"{OUT}/Lamina_Pericial_P01_Valle_Escondido_LT220_Caone.pdf")
    fig.savefig("/tmp/qa/P01.png", dpi=300)
    plt.close(fig)
    return esc_txt


def lamina_p02(lotes, base, faja, torres, caminos):
    fig = plt.figure(figsize=(16.54, 11.69))
    marco(fig, "P-02 · DETALLE DE LOS 14 LOTES DIRECTAMENTE AFECTADOS",
          "Proyecto Valle Escondido / Fundo Hualañé  ·  Línea Eléctrica 1x220 kV Caone–Mataquito\n"
          "Hualañé – Región del Maule  ·  Folio " + FOLIO + "-P02")

    ax = fig.add_axes([0.032, 0.045, 0.615, 0.855])
    sel = [l for l in lotes if l["rol"] == RECL and l["n"] in PES]
    u = unary_union([l["poly"] for l in sel])
    fx0, fy0, fx1, fy1 = u.bounds
    cx, cy = (fx0 + fx1) / 2, (fy0 + fy1) / 2
    w = max(fx1 - fx0, fy1 - fy0) * 1.25
    ratio = 0.615 * 16.54 / (0.855 * 11.69)
    bbox = (cx - w * ratio / 2, cy - w / 2, cx + w * ratio / 2, cy + w / 2)
    faja_loc = faja.intersection(u.buffer(220))
    dibujar_mapa(ax, sel, base, faja_loc, torres, caminos, bbox,
                 etiquetas={(RECL, n) for n in PES}, label_pct=True)
    ancho_m = bbox[2] - bbox[0]
    esc = escala_real(0.615 * 16.54, ancho_m)
    barra_escala(ax, bbox[0] + ancho_m * 0.04, bbox[1] + ancho_m * 0.045, ancho_m)
    norte(ax, bbox[2] - ancho_m * 0.05, bbox[1] + ancho_m * 0.04, ancho_m * 0.06)
    ax.set_xlabel("Este UTM 19S (m) – EPSG:32719", fontsize=5.8)
    ax.set_ylabel("Norte UTM 19S (m) – EPSG:32719", fontsize=5.8)
    esc_txt = f"1:{round(esc / 100) * 100:,}".replace(",", ".") + " (A3)"

    p = fig.add_axes([0.663, 0.045, 0.305, 0.855])
    p.axis("off")
    y = 1.0
    y = caja(p, y, 0.215, "CUADRO DE LOS 14 LOTES NOTIFICADOS")
    cols = [("Lote", 0.02, "l"), ("Sup. m²", 0.235, "r"), ("PES m²", 0.42, "r"),
            ("% af.", 0.575, "r"), ("Rem. m²", 0.755, "r"), ("Clase", 0.98, "r")]
    for t, xx, al in cols:
        p.text(xx, y, t, transform=p.transAxes, fontsize=5.2, fontweight="bold", va="top",
               ha="left" if al == "l" else "right")
    y -= 0.0125
    for n, (s, pes, pct) in sorted(PES.items(), key=lambda kv2: -kv2[1][2]):
        cl = base[(RECL, n)]["clase"]
        destacado = n in (61, 60, 57, 56, 55)
        fw = "bold" if destacado else "normal"
        vals = [(str(n), 0.02, "l"), (f"{s:,}".replace(",", "."), 0.235, "r"),
                (f"{pes:,}".replace(",", "."), 0.42, "r"),
                (f"{pct:.1f}".replace(".", ",") + " %", 0.575, "r"),
                (f"{s - pes:,}".replace(",", "."), 0.755, "r"), (cl, 0.98, "r")]
        for t, xx, al in vals:
            p.text(xx, y, t, transform=p.transAxes, fontsize=5.2, va="top",
                   ha="left" if al == "l" else "right", fontweight=fw)
        y -= 0.0105
    p.plot([0.02, 0.98], [y + 0.006, y + 0.006], transform=p.transAxes, color="black", lw=0.5)
    y -= 0.002
    tot = [("TOTAL", 0.02, "l"), ("76.700", 0.235, "r"), ("35.777", 0.42, "r"),
           ("46,6 %", 0.575, "r"), ("40.923", 0.755, "r"), ("—", 0.98, "r")]
    for t, xx, al in tot:
        p.text(xx, y, t, transform=p.transAxes, fontsize=5.2, va="top", fontweight="bold",
               ha="left" if al == "l" else "right")
    y -= 0.0135
    y = parrafo(p, y, "En negrita los lotes 61, 60, 57, 56 y 55: mayor porcentaje de\n"
                      "afectación física por la faja PES notificada.", dy=0.0105)
    y -= 0.018

    y = caja(p, y, 0.062, "SUPERFICIE REMANENTE Y LIBRE")
    y = kv(p, y, "Superficie notificada total", "76.700 m²", dy=0.0135)
    y = kv(p, y, "Superficie gravada por PES", "35.777 m²", dy=0.0135)
    y = kv(p, y, "Superficie remanente libre", "40.923 m²", dy=0.0135, bold=True)
    y -= 0.014

    y = simbologia(p, y)

    y = caja(p, y, 0.070, "NOTA")
    y = parrafo(p, y,
                "El hachurado representa la afectación física (faja PES). El color de\n"
                "fondo del polígono representa la clase de exposición visual V5 (V0–V4).\n"
                "Ambos conceptos son independientes y no deben sumarse ni contabilizarse\n"
                "dos veces. Los porcentajes rotulados corresponden a los PES oficiales\n"
                "notificados y no a la reconstrucción GIS.", dy=0.0105)
    y -= 0.014
    cartucho(p, y, "P02", esc_txt)

    fig.savefig(f"{OUT}/Lamina_Pericial_P02_14_Lotes_Afectados_LT220_Caone.pdf")
    fig.savefig("/tmp/qa/P02.png", dpi=300)
    plt.close(fig)


# ----------------------------------------------------------------- tablas --
def exportar_tablas(lotes, base, faja):
    hdr = ["Rol", "Lote", "Superficie GIS m2", "Superficie PES oficial m2",
           "% afectacion PES", "IDV", "Clase V", "Dist. min. faja (m)",
           "Dist. min. estructura (m)", "Cruzado por faja (GIS)",
           "Sup. faja en lote GIS (m2)", "Centroide X UTM19S", "Centroide Y UTM19S",
           "Lote notificado PES"]
    rows = []
    for l in sorted(lotes, key=lambda l: (l["rol"], l["n"])):
        b = base[(l["rol"], l["n"])]
        notif = l["rol"] == RECL and l["n"] in PES
        rows.append([l["rol"], l["n"], l["area"],
                     PES[l["n"]][1] if notif else "",
                     PES[l["n"]][2] if notif else "",
                     b["idv"], b["clase"], b["d_faja"], b["d_torre"],
                     "Si" if l["cruzada"] else "No", l["area_faja"], b["x"], b["y"],
                     "Si" if notif else "No"])
    wb = Workbook()
    ws = wb.active
    ws.title = "LOTES_REPRESENTADOS"
    ws.append(hdr)
    for r in rows:
        ws.append(r)
    for i, w in enumerate([13, 7, 16, 22, 16, 9, 9, 17, 21, 19, 22, 18, 18, 18], 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    ws.freeze_panes = "A2"
    wb.save(f"{OUT}/Lamina_Pericial_V5_Lotes_Representados.xlsx")
    with open(f"{OUT}/Lamina_Pericial_V5_Lotes_Representados.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(hdr)
        w.writerows(rows)

    T4326 = Transformer.from_crs("EPSG:32719", "EPSG:4326", always_xy=True)
    feats = []
    for l in lotes:
        b = base[(l["rol"], l["n"])]
        notif = l["rol"] == RECL and l["n"] in PES
        feats.append({"type": "Feature", "geometry": mapping(l["poly"]), "properties": {
            "rol": l["rol"], "lote": l["n"], "area_gis_m2": l["area"], "idv": b["idv"],
            "clase_v": b["clase"], "cruzado_faja": bool(l["cruzada"]),
            "sup_faja_gis_m2": l["area_faja"],
            "pes_oficial_m2": PES[l["n"]][1] if notif else None,
            "pct_pes": PES[l["n"]][2] if notif else None}})
    feats.append({"type": "Feature", "geometry": mapping(faja),
                  "properties": {"capa": "Faja de servidumbre PES (reconstrucción GIS V5)"}})
    gj = {"type": "FeatureCollection",
          "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::32719"}},
          "features": feats}
    json.dump(gj, open(f"{OUT}/Lamina_Pericial_V5_Poligonos_EPSG32719.geojson", "w"))
    return len(rows)


def main():
    os.makedirs("/tmp/qa", exist_ok=True)
    lotes, faja, torres = load_geo()
    base = load_base()
    caminos = load_caminos()
    errs, cnt = validar(lotes, base, torres)
    if errs:
        print("VALIDACIÓN FALLIDA — EXPORTACIÓN DETENIDA:")
        for e in errs:
            print("  -", e)
        raise SystemExit(1)
    print("Validaciones 1-11: OK  ·", cnt, "· torres", len(torres), "· caminos", len(caminos))
    esc = lamina_p01(lotes, base, faja, torres, caminos, cnt)
    lamina_p02(lotes, base, faja, torres, caminos)
    n = exportar_tablas(lotes, base, faja)
    print("Láminas generadas. Escala P-01", esc, "· filas tabla", n)


if __name__ == "__main__":
    main()
